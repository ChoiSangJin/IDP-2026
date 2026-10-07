import io
import json
from xml.etree import ElementTree
import pytest
from openpyxl import load_workbook
from auto_dd.sample import sample
from auto_dd.model import validate
from auto_dd.export import export, layout
from auto_dd.web import create_app


def test_roundtrip_and_editor_layout():
    data = sample()
    data["tables"][0]["position"] = {"x": 620, "y": 350}
    data["tables"][0]["description"] = "수정한 설명"
    restored = validate(json.loads(export(data, "json")))
    assert restored == data
    assert layout(restored)["public.customers"] == {"x": 620, "y": 350}
    assert restored["tables"][1]["foreign_keys"][0]["target_table"] == "public.customers"


def test_excel_has_metadata_relationships_and_literal_formulas():
    data = sample()
    data["tables"][0]["description"] = '=HYPERLINK("https://example.com")'
    wb = load_workbook(io.BytesIO(export(data, "xlsx")))
    assert wb.sheetnames == ["Tables", "Columns", "Relationships"]
    assert wb["Tables"].max_row == 4
    assert wb["Columns"].max_row == 12
    assert wb["Relationships"]["D2"].value == "public.customers"
    assert wb["Tables"]["C2"].data_type == "s"
    assert wb["Tables"]["C2"].value.startswith("=HYPERLINK")


def test_svg_escapes_untrusted_metadata_and_draws_relationships():
    data = sample()
    data["tables"][0]["name"] = '<script>alert("x")</script>'
    root = ElementTree.fromstring(export(data, "svg"))
    ns = {"s": "http://www.w3.org/2000/svg"}
    assert not root.findall(".//s:script", ns)
    assert len(root.findall("s:path", ns)) == 2
    assert '<script>alert("x")</script>' in ''.join(root.itertext())


def test_mermaid_uses_safe_ids_and_composite_fk():
    result = export(sample(), "mmd").decode()
    assert 'T0["public.customers"]' in result
    assert 'T0 ||..o{ T1 : "orders_customer_fk"' in result


@pytest.mark.parametrize("mutation", [
    lambda d: d["tables"].append(d["tables"][0]),
    lambda d: d["tables"][0]["columns"].append(d["tables"][0]["columns"][0]),
    lambda d: d["tables"][0].update(position={"x": float("nan"), "y": 0}),
    lambda d: d["tables"][1]["foreign_keys"][0].update(columns=["missing"]),
    lambda d: d["tables"][1]["foreign_keys"][0].update(target_columns=["missing"]),
    lambda d: d["tables"][0]["columns"][0].update(nullable="yes"),
])
def test_invalid_catalog_rejected(mutation):
    data = sample()
    mutation(data)
    with pytest.raises(ValueError):
        validate(data)


def test_http_import_export_auth_and_rebinding():
    app = create_app("test-token")
    client = app.test_client()
    headers = {"X-Auto-DD-Token": "test-token"}
    assert client.get("/health").json["status"] == "ok"
    assert client.get("/").status_code == 200
    assert client.get("/assets/app.js").status_code == 200
    assert client.get("/api/sample").status_code == 403
    assert client.get("/api/sample", headers=headers, base_url="http://evil.example").status_code == 403
    assert client.get("/api/sample", headers={**headers, "Origin": "https://evil.example"}).status_code == 403
    data = client.get("/api/sample", headers=headers).json
    data["tables"][0]["description"] = "한글 설명"
    assert client.post("/api/validate", headers=headers, json=data).json["tables"][0]["description"] == "한글 설명"
    for kind in ("xlsx", "json", "svg", "mmd"):
        response = client.post("/api/export/"+kind, headers=headers, json=data)
        assert response.status_code == 200
        assert len(response.data) > 100
        assert response.headers["Content-Disposition"].startswith("attachment;")
    assert client.post("/api/validate", headers=headers, json={}).status_code == 400
    assert client.post("/api/export/bad", headers=headers, json=data).status_code == 400


def test_connection_errors_do_not_disclose_credentials(monkeypatch):
    import auto_dd.web as web
    def fail(*args, **kwargs):
        raise RuntimeError("password=private-test-password")
    monkeypatch.setattr(web, "postgres", fail)
    response = create_app("test-token").test_client().post("/api/extract", headers={"X-Auto-DD-Token": "test-token"}, json={"source": "postgresql", "dsn": "password=private-test-password"})
    assert response.status_code == 502
    assert b"private-test-password" not in response.data
