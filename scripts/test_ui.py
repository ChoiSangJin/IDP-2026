"""Optional browser integration checks (pip install playwright && playwright install chromium)."""
import io
import json
import os
import threading
from pathlib import Path
from tempfile import TemporaryDirectory
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
from werkzeug.serving import make_server
from auto_dd.web import create_app

app = create_app()
server = make_server("127.0.0.1", 0, app, threaded=True)
thread = threading.Thread(target=server.serve_forever, daemon=True)
thread.start()
url = f"http://127.0.0.1:{server.server_port}/#token={app.config['TOKEN']}"
try:
    with sync_playwright() as p, TemporaryDirectory() as temp:
        options = {"headless": True}
        if os.environ.get("AUTO_DD_CHROMIUM"):
            options["executable_path"] = os.environ["AUTO_DD_CHROMIUM"]
        browser = p.chromium.launch(**options)
        page = browser.new_page(viewport={"width": 1440, "height": 1000})
        errors, external = [], []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.on("request", lambda request: external.append(request.url) if not request.url.startswith("http://127.0.0.1:") else None)
        page.goto(url)
        expect(page.locator("#empty")).to_be_visible()
        page.locator("#empty-sample").click()
        expect(page.locator("#table-count")).to_have_text("3")
        page.locator(".table-card .table-actions button").first.click()
        page.locator("#edit-description").fill("브라우저에서 수정한 설명")
        page.locator("#table-form button[type=submit]").click()
        expect(page.locator("#table-dialog")).not_to_be_visible()
        expect(page.locator(".table-card").first).to_contain_text("브라우저에서 수정한 설명")
        page.locator("#new-table").click()
        page.locator("#edit-name").fill("loyalty")
        page.locator("#new-column").click()
        page.locator("#column-editor [data-field=name]").nth(1).fill("customer_id")
        page.locator("#column-editor [data-field=type]").nth(1).fill("bigint")
        page.locator("#table-form button[type=submit]").click()
        expect(page.locator("#table-count")).to_have_text("4")
        page.locator(".tabs [data-view=erd]").click()
        page.locator("#add-relation").click()
        page.locator("#fk-name").fill("loyalty_customer_fk")
        page.locator("#fk-source").select_option("public.loyalty")
        page.locator("#fk-target").select_option("public.customers")
        page.locator("#fk-columns").fill("customer_id")
        page.locator("#fk-target-columns").fill("id")
        page.locator("#relation-form button[type=submit]").click()
        expect(page.locator("#relation-count")).to_have_text("3")
        node = page.locator(".node-head").first
        rect = node.bounding_box()
        page.mouse.move(rect["x"]+30, rect["y"]+15)
        page.mouse.down()
        page.mouse.move(rect["x"]+130, rect["y"]+95, steps=10)
        page.mouse.up()
        page.locator("#format").select_option("json")
        with page.expect_download() as download:
            page.locator("#export").click()
        path = Path(temp)/"catalog.json"
        download.value.save_as(path)
        data = json.loads(path.read_text())
        assert data["tables"][0]["position"] == {"x": 140, "y": 120}
        assert data["tables"][-1]["foreign_keys"][0]["target_table"] == "public.customers"
        page.locator("#file").set_input_files({"name":"invalid.json","mimeType":"application/json","buffer":b'{"version":9}'})
        expect(page.locator("#notice")).to_contain_text("version: 1")
        expect(page.locator("#table-count")).to_have_text("4")
        page.locator("#file").set_input_files(path)
        expect(page.locator("#notice")).to_contain_text("JSON 카탈로그를 가져왔습니다")
        page.locator("#format").select_option("xlsx")
        with page.expect_download() as download:
            page.locator("#export").click()
        excel = Path(temp)/"spec.xlsx"
        download.value.save_as(excel)
        assert load_workbook(excel)["Tables"].max_row == 5
        page.locator(".tabs [data-view=catalog]").click()
        page.locator("#search").fill("loyalty")
        expect(page.locator(".table-card")).to_have_count(1)
        page.locator("#search").fill("")
        page.set_viewport_size({"width": 390, "height": 844})
        expect(page.locator("#new-table")).to_be_visible()
        assert not errors, errors
        assert not external, external
        browser.close()
        print("Browser checks passed: edit, create, relationships, drag, JSON roundtrip, invalid import, Excel, search, mobile, no external requests")
finally:
    server.shutdown()
    server.server_close()
